"""
Anode Tracking Report checks.

Runs the 7 data-quality checks described in the requirements sheet against
the "Details" tab of an Anode Tracking Report workbook:

  1. powderName blank                              -> list rows
  2. quantity below 4000                           -> list rows
  3. category == 'B' and partNumber's last 4 chars
     start with '76'                               -> daily count/list summary
  4. length/thickness/wiresize == the target
     dimension set (0.054 / 0.041 / 0.0118)         -> daily count/qty summary,
                                                        flagged if total qty > 250,000
                                                        (width is not part of the match)
  5. duplicate anodeLotID within the report         -> list rows
  6. anodeLotID also present, with a non-blank
     PlanDate, in the master "Anode 3-4-8.19.xlsx"
     workbook's "details" sheet (LOT_NUMBER column)   -> list rows
     (a trailing letter on the anodeLotID is
     stripped before matching, and the match is by
     substring containment rather than exact equality)
  7. component column, characters 11-12 == '75' or
     '63', and subInventory is Intransit or
     ANODE-INSP                                     -> list rows

Results are written to a multi-sheet Excel report.
"""

import argparse
from collections import defaultdict
from datetime import datetime

import openpyxl
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

SHEET_NAME = "Details"
MASTER_SHEET_NAME = "details"
QTY_THRESHOLD = 4000
TARGET_DIMS = {"length": 0.054, "thickness": 0.041, "wiresize": 0.0118}
DAILY_QTY_LIMIT = 250_000
FLAGGED_C_CODES = {"75", "63"}
FLAGGED_SUBINVENTORY = {"Intransit", "ANODE-INSP"}

EXCEEDS_250K_FILL = PatternFill(start_color="FF92D050", end_color="FF92D050", fill_type="solid")

EXPORT_COLUMNS = [
    "anodeLotID", "partNumber", "component", "powderName", "planDate", "quantity",
    "category", "width", "length", "thickness", "wiresize", "subInventory",
]


HEADER_ALIASES = {"Anode": "component"}  # column C is labeled differently across file versions


def load_rows(path, sheet_name=SHEET_NAME):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet_name]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    headers = [HEADER_ALIASES.get(h, h) for h in headers]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def load_master_planned_lots(path, sheet_name=MASTER_SHEET_NAME):
    """LOT_NUMBER values from the master workbook's rows where PlanDate is NOT blank."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    headers = next(rows_iter)
    lot_idx = headers.index("LOT_NUMBER")
    date_idx = headers.index("PlanDate")
    lots = []
    for r in rows_iter:
        plan_date = r[date_idx]
        if not is_blank(plan_date):
            lot = r[lot_idx]
            if lot:
                lots.append(lot)
    return lots


def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")


def to_date(value):
    return value.date() if isinstance(value, datetime) else value


# --- checks -----------------------------------------------------------------

def check_blank_powder_name(rows):
    return [r for r in rows if is_blank(r.get("powderName"))]


def check_low_quantity(rows, threshold=QTY_THRESHOLD):
    return [r for r in rows if isinstance(r.get("quantity"), (int, float)) and r["quantity"] < threshold]


def check_category_b_partnumber_76(rows):
    out = []
    for r in rows:
        if r.get("category") != "B":
            continue
        pn = r.get("partNumber")
        if isinstance(pn, str) and len(pn) >= 4 and pn[-4:-2] == "76":
            out.append(r)
    return out


def check_target_dimensions(rows, dims=TARGET_DIMS, tol=1e-6):
    out = []
    for r in rows:
        if all(r.get(k) is not None and abs(r[k] - v) < tol for k, v in dims.items()):
            out.append(r)
    return out


def check_duplicate_lot_ids(rows):
    seen = defaultdict(list)
    for r in rows:
        lot = r.get("anodeLotID")
        if lot:
            seen[lot].append(r)
    return {lot: items for lot, items in seen.items() if len(items) > 1}


def strip_trailing_letter(lot_id):
    if isinstance(lot_id, str) and lot_id and lot_id[-1].isalpha():
        return lot_id[:-1]
    return lot_id


def check_cross_file_duplicate_lot_ids(rows_current, master_lots):
    master_lots = [m for m in master_lots if isinstance(m, str)]
    out = []
    for r in rows_current:
        lot = r.get("anodeLotID")
        if not lot:
            continue
        norm = strip_trailing_letter(lot)
        if any(norm in m or m in norm for m in master_lots):
            out.append(r)
    return out


def check_c_pos12_status(rows):
    out = []
    for r in rows:
        c_val = r.get("component")
        if isinstance(c_val, str) and len(c_val) >= 12 and c_val[10:12] in FLAGGED_C_CODES:
            if r.get("subInventory") in FLAGGED_SUBINVENTORY:
                out.append(r)
    return out


def daily_summary(rows, date_field="planDate"):
    buckets = defaultdict(list)
    for r in rows:
        buckets[to_date(r.get(date_field))].append(r)
    summary = []
    for d, items in sorted(buckets.items(), key=lambda kv: (kv[0] is None, kv[0])):
        total_qty = sum(i.get("quantity") or 0 for i in items)
        summary.append({
            "date": d,
            "count": len(items),
            "total_quantity": total_qty,
            "exceeds_250k": total_qty > DAILY_QTY_LIMIT,
            "rows": items,
        })
    return summary


# --- report writing -----------------------------------------------------------

def write_detail_sheet(wb, title, rows, columns=EXPORT_COLUMNS):
    ws = wb.create_sheet(title)
    ws.append(columns)
    for r in rows:
        ws.append([r.get(c) for c in columns])
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(col) + 2)
    return ws


def write_duplicate_lot_sheet(wb, title, dup_map):
    ws = wb.create_sheet(title)
    ws.append(["anodeLotID", "occurrences"] + EXPORT_COLUMNS)
    for lot, items in dup_map.items():
        for r in items:
            ws.append([lot, len(items)] + [r.get(c) for c in EXPORT_COLUMNS])
    return ws


def write_daily_summary_sheets(wb, prefix, summary):
    overview = wb.create_sheet(f"{prefix}_daily")
    overview.append(["date", "count", "total_quantity", "exceeds_250K"])
    for s in summary:
        overview.append([s["date"], s["count"], s["total_quantity"], s["exceeds_250k"]])
        if s["exceeds_250k"]:
            overview.cell(row=overview.max_row, column=4).fill = EXCEEDS_250K_FILL

    detail = wb.create_sheet(f"{prefix}_detail")
    detail.append(["date"] + EXPORT_COLUMNS)
    for s in summary:
        for r in s["rows"]:
            detail.append([s["date"]] + [r.get(c) for c in EXPORT_COLUMNS])


def build_report(current_rows, master_lots, output_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    write_detail_sheet(wb, "1_BlankPowderName", check_blank_powder_name(current_rows))
    write_detail_sheet(wb, "2_LowQuantity_lt4000", check_low_quantity(current_rows))

    cat_b_76 = check_category_b_partnumber_76(current_rows)
    write_daily_summary_sheets(wb, "3_CatB_PN76", daily_summary(cat_b_76))

    target_dim_rows = check_target_dimensions(current_rows)
    write_daily_summary_sheets(wb, "4_TargetDims", daily_summary(target_dim_rows))

    write_duplicate_lot_sheet(wb, "5_DuplicateLotID", check_duplicate_lot_ids(current_rows))
    write_detail_sheet(
        wb, "6_CrossFileDupLotID",
        check_cross_file_duplicate_lot_ids(current_rows, master_lots),
    )
    write_detail_sheet(wb, "7_C12_IntransitOrInsp", check_c_pos12_status(current_rows))

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description="Run Anode Tracking Report checks.")
    parser.add_argument("--current", default="Anode Tracking Report Aug.19th.xlsx")
    parser.add_argument("--master", default="Anode 3-4-8.19.xlsx")
    parser.add_argument("--output", default="Anode_Tracking_Check_Report.xlsx")
    parser.add_argument("--sheet", default=SHEET_NAME)
    parser.add_argument("--master-sheet", default=MASTER_SHEET_NAME)
    args = parser.parse_args()

    current_rows = load_rows(args.current, args.sheet)
    master_lots = load_master_planned_lots(args.master, args.master_sheet)

    build_report(current_rows, master_lots, args.output)

    print(f"Loaded {len(current_rows)} rows from '{args.current}' ({args.sheet}).")
    print(f"Loaded {len(master_lots)} non-blank-PlanDate LOT_NUMBERs from '{args.master}' ({args.master_sheet}).")
    print(f"1. Blank powderName: {len(check_blank_powder_name(current_rows))}")
    print(f"2. Quantity < {QTY_THRESHOLD}: {len(check_low_quantity(current_rows))}")
    print(f"3. Category B & partNumber 76xx: {len(check_category_b_partnumber_76(current_rows))}")
    print(f"4. Target dimension rows: {len(check_target_dimensions(current_rows))}")
    dup = check_duplicate_lot_ids(current_rows)
    print(f"5. Duplicate LotIDs: {len(dup)} lot(s), {sum(len(v) for v in dup.values())} row(s)")
    print(f"6. LotIDs also planned (non-blank PlanDate) in master: {len(check_cross_file_duplicate_lot_ids(current_rows, master_lots))}")
    print(f"7. C[11:12] in 75/63 & subInventory Intransit/ANODE-INSP: {len(check_c_pos12_status(current_rows))}")
    print(f"Report written to '{args.output}'.")


if __name__ == "__main__":
    main()
