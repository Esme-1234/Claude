"""
Convert CategorySIPMaintain.xlsx (sheets: 687, D, Rest) into a long-format
upload sheet: UploadType / PartNumber / PlanDate / BatchesPerDay.

Scope: one week, starting from the date in the Summary sheet's D1 cell
(the most recent Saturday) through the following Friday.

Extraction rule: a cell counts as a genuine "batches per day" entry only if
it is a literal number typed directly into the sheet (not a formula). This
naturally includes every hand-entered part-number row -- whether it sits in
a QTY-labeled block (687/D sheets) or stands alone further down the sheet
(e.g. rows 33+ on the D sheet) -- and naturally excludes every derived cell:
the QTY row itself (QTY = part row * batch size), and Total/Non-687/subtotal
rows (all SUM-style formulas), since those are formulas rather than literals.

Usage:
    python3 build_long_format.py CategorySIPMaintain.7.23.xlsx CategorySIPMaintain.7.23.long.xlsx
"""
import sys
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill


def get_start_date(wbd):
    """Week start date is read from the Summary sheet's D1 cell."""
    v = wbd['Summary']['D1'].value
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    raise ValueError(f"Summary!D1 is not a date: {v!r}")


def find_target_cols(ws, target_dates):
    """Locate the columns in row 1 whose date falls in target_dates."""
    cols = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if isinstance(v, datetime.datetime) and v.date() in target_dates:
            cols.append(col)
    cols.sort()
    return cols


def extract_sheet(wsf, wsd, sheet_name, target_dates, results):
    cols = find_target_cols(wsf, target_dates)
    assert len(cols) == 7, (sheet_name, cols)

    for r in range(2, wsf.max_row + 1):
        part = wsf.cell(row=r, column=1).value
        if part is None or part == 'QTY':
            continue
        for col, d in zip(cols, target_dates):
            raw = wsf.cell(row=r, column=col).value
            if isinstance(raw, str) and raw.startswith('='):
                continue  # derived cell (QTY multiplier, Total/subtotal roll-up) -- skip
            if isinstance(raw, (int, float)) and not isinstance(raw, bool) and raw != 0:
                results.append((part, d, wsd.cell(row=r, column=col).value))


def build(src_path, out_path):
    wbf = openpyxl.load_workbook(src_path, data_only=False)
    wbd = openpyxl.load_workbook(src_path, data_only=True)

    start_date = get_start_date(wbd)
    target_dates = [start_date + datetime.timedelta(days=i) for i in range(7)]
    weekend_dates = {target_dates[0], target_dates[1]}

    results = []  # (part_number, date, batches_per_day)
    for sheet_name in ('687', 'D', 'Rest'):
        extract_sheet(wbf[sheet_name], wbd[sheet_name], sheet_name, target_dates, results)

    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = 'Sheet1'

    header_fill = PatternFill(start_color='FFD9D9D9', end_color='FFD9D9D9', fill_type='solid')
    weekend_fill = PatternFill(start_color='FFFCE4D6', end_color='FFFCE4D6', fill_type='solid')
    font = Font(name='Arial', size=10)
    bold_font = Font(name='Arial', size=10, bold=True)

    for c, h in enumerate(['UploadType', 'PartNumber', 'PlanDate', 'BatchesPerDay'], start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = bold_font
        cell.fill = header_fill

    r = 2
    for part, d, v in results:
        ws.cell(row=r, column=1, value='ADD').font = font
        ws.cell(row=r, column=2, value=part).font = font
        dcell = ws.cell(row=r, column=3, value=datetime.datetime(d.year, d.month, d.day))
        dcell.font = font
        dcell.number_format = 'm/d/yyyy'
        ws.cell(row=r, column=4, value=v).font = font
        if d in weekend_dates:
            for c in range(1, 5):
                ws.cell(row=r, column=c).fill = weekend_fill
        r += 1

    for col, width in {'A': 12, 'B': 26, 'C': 12, 'D': 15}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    out_wb.save(out_path)
    print(f'Wrote {r - 2} rows to {out_path}')


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'CategorySIPMaintain.7.23.xlsx'
    out = sys.argv[2] if len(sys.argv) > 2 else 'CategorySIPMaintain.7.23.long.xlsx'
    build(src, out)
