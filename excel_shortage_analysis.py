"""
Build a shortage report from three planning Excel files:

1. planningWarningMessage.xlsx - Sheet1 has one row per planning warning.
   Rows whose warningMessage starts with "CalculatedStartDate" identify the
   partNumbers that need a shortage check.
2. Loading *.xlsx - two sheets are used:
   - "Summary" is a pivot export keyed by PARENT_ITEM, with one column per
     plan week (header is the week number, e.g. 202628). The first such
     week column gives the required quantity for a partNumber.
   - "KEMET_GOP_ATO_ANALYSIS_DAILY" has a PARENT_ITEM column and an
     "Anode Category" column; this gives the category for each
     partNumber.
3. PlanningResult.xlsx - Sheet1 has one row per planned lot (partNumber,
   planDate, quantity). matchedQty is the sum of quantity over every row
   whose partNumber matches - the total already covered by planning.

For every target partNumber: gap = matchedQty - requiredQty. A negative gap
means planning hasn't covered the requirement yet (a shortage). The output
workbook has "AllPartNumbers" (every target partNumber) and "ShortageOnly"
(gap < 0 rows only - the answer to "how much quantity is short").

Usage:
    python excel_shortage_analysis.py \
        --warning "planningWarningMessage (2).xlsx" \
        --loading "Loading July.8th-New.xlsx" \
        --planning-result "PlanningResult (7).xlsx" \
        --output shortage_report.xlsx
"""

import argparse
import sys
from collections import defaultdict

import openpyxl


def _find_header_row(ws, required_header, max_scan_rows=15, max_col=40):
    for r in range(1, max_scan_rows + 1):
        values = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        if required_header in values:
            return r, values
    raise ValueError(f"Could not find a header cell '{required_header}' in the first {max_scan_rows} rows")


def get_target_part_numbers(warning_path, sheet_name=None, prefix="CalculatedStartDate"):
    wb = openpyxl.load_workbook(warning_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    header_row, header = _find_header_row(ws, "warningMessage", max_scan_rows=1)
    pn_col = header.index("partNumber") + 1
    wm_col = header.index("warningMessage") + 1

    seen, seen_set = [], set()
    for row in ws.iter_rows(min_row=header_row + 1):
        wm = row[wm_col - 1].value
        pn = row[pn_col - 1].value
        if wm and pn and str(wm).startswith(prefix) and pn not in seen_set:
            seen.append(pn)
            seen_set.add(pn)
    return seen


def get_required_quantities(loading_path, sheet_name="Summary"):
    wb = openpyxl.load_workbook(loading_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header_row, header = _find_header_row(ws, "PARENT_ITEM", max_scan_rows=10)
    parent_col = header.index("PARENT_ITEM")
    def _looks_like_week(value):
        return isinstance(value, int) or (isinstance(value, str) and value.isdigit())

    week_col = next(i for i in range(parent_col + 1, len(header)) if _looks_like_week(header[i]))
    week_label = header[week_col]

    required = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        pn = row[parent_col]
        if pn:
            required[pn] = row[week_col] or 0
    return required, week_label


def get_matched_quantities(planning_result_path, sheet_name=None):
    wb = openpyxl.load_workbook(planning_result_path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    header_row, header = _find_header_row(ws, "partNumber", max_scan_rows=1)
    pn_idx = header.index("partNumber")
    qty_idx = header.index("quantity")

    matched = defaultdict(float)
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        pn = row[pn_idx]
        if pn:
            matched[pn] += row[qty_idx] or 0
    return matched


def get_categories(loading_path, sheet_name="KEMET_GOP_ATO_ANALYSIS_DAILY"):
    wb = openpyxl.load_workbook(loading_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    header_row, header = _find_header_row(ws, "PARENT_ITEM", max_scan_rows=15)
    parent_col = header.index("PARENT_ITEM")
    cat_col = header.index("Anode Category")

    categories = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        pn = row[parent_col]
        if pn and pn not in categories:
            categories[pn] = row[cat_col]
    return categories


def build_report(warning_path, loading_path, planning_result_path):
    part_numbers = get_target_part_numbers(warning_path)
    required, week_label = get_required_quantities(loading_path)
    matched = get_matched_quantities(planning_result_path)
    categories = get_categories(loading_path)

    rows = []
    for pn in part_numbers:
        req = required.get(pn, 0)
        mat = matched.get(pn, 0)
        rows.append((categories.get(pn, ""), pn, req, mat, mat - req))
    return rows, week_label


def save(rows, week_label, output_path):
    wb = openpyxl.Workbook()
    header = ["category", "partNumber", week_label, "matchedQty", "gap"]

    ws_all = wb.active
    ws_all.title = "AllPartNumbers"
    ws_all.append(header)
    for row in rows:
        ws_all.append(list(row))

    ws_short = wb.create_sheet("ShortageOnly")
    ws_short.append(header)
    for row in rows:
        if row[4] < 0:
            ws_short.append(list(row))

    for ws in (ws_all, ws_short):
        for col, width in zip("ABCDE", (14, 28, 14, 14, 16)):
            ws.column_dimensions[col].width = width
        for row in range(2, ws.max_row + 1):
            for col in "CDE":
                ws[f"{col}{row}"].number_format = "#,##0;(#,##0)"

    wb.save(output_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--warning", required=True, help="Path to the planningWarningMessage .xlsx file")
    parser.add_argument("--loading", required=True, help="Path to the Loading .xlsx file (Summary sheet)")
    parser.add_argument("--planning-result", required=True, help="Path to the PlanningResult .xlsx file")
    parser.add_argument("--output", required=True, help="Path to write the shortage report .xlsx file")
    args = parser.parse_args()

    rows, week_label = build_report(args.warning, args.loading, args.planning_result)
    save(rows, week_label, args.output)

    shortage_count = sum(1 for r in rows if r[4] < 0)
    print(f"{len(rows)} target partNumbers, {shortage_count} with a shortage -> {args.output}")


if __name__ == "__main__":
    sys.exit(main())
